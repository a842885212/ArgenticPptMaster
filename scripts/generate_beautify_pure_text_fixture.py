#!/usr/bin/env python3
"""Generate the pure-text PPTX fixture used by beautify stage-1 tests."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


SLIDES = [
    ["Pure Text Beautify", "Every slide remains in order."],
    ["Repeated text", "Keep me", "Keep me", "Dense page starts here"],
    ["Dense page", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight"],
    ["Pause", ""],
]


def main() -> None:
    output = Path("src/test/resources/beautify-pptx/pure-text-mvp.pptx")
    presentation = Presentation()
    layout = presentation.slide_layouts[6]
    for lines in SLIDES:
        slide = presentation.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(5.8))
        frame = box.text_frame
        frame.clear()
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(30 if index == 0 else 18)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    print(output)


if __name__ == "__main__":
    main()
