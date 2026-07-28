#!/usr/bin/env python3
"""Generate the synthetic PPTX fixture used by the beautify-pptx phase-0 baseline."""

from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
TITLE_COLOR = "16324F"
ACCENT_COLOR = "2878B5"
TEXT_COLOR = "263238"
MUTED_COLOR = "5C6B73"


def add_text(slide, text: str, left: float, top: float, width: float, height: float,
             size: float, color: str = TEXT_COLOR, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)


def add_title(slide, title: str) -> None:
    add_text(slide, title, 0.7, 0.45, 11.8, 0.55, 26, TITLE_COLOR, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.18), Inches(2.0), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor.from_string(ACCENT_COLOR)
    line.line.fill.background()


def add_bullets(slide, bullets: list[str], top: float = 1.65, size: float = 20) -> None:
    box = slide.shapes.add_textbox(Inches(1.0), Inches(top), Inches(10.8), Inches(5.0))
    frame = box.text_frame
    frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
        paragraph.space_after = Pt(12)


def make_image() -> Path:
    image_path = Path(tempfile.mkstemp(prefix="beautify-phase0-", suffix=".png")[1])
    image = Image.new("RGB", (960, 540), "#16324F")
    draw = ImageDraw.Draw(image)
    for offset in range(0, 960, 24):
        shade = 70 + offset // 12
        draw.rectangle((offset, 0, offset + 24, 540), fill=(24, min(shade, 150), 180))
    draw.rounded_rectangle((130, 120, 830, 420), radius=32, fill="#F3F7FA")
    draw.text((245, 230), "SYNTHETIC IMAGE", fill="#16324F", font_size=48)
    image.save(image_path)
    return image_path


def build_presentation(output: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    add_text(slide, "Beautify Phase 0 Baseline", 0.8, 1.8, 11.4, 0.8, 34, TITLE_COLOR, True)
    add_text(slide, "Synthetic deck — no external data.", 0.85, 2.85, 9.5, 0.45, 20, MUTED_COLOR)
    add_text(slide, "Strict 1:1 content and data validation", 0.85, 4.9, 7.0, 0.4, 18, ACCENT_COLOR, True)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Text Fidelity")
    add_text(slide, "Every sentence remains on its source slide.", 1.0, 1.65, 10.5, 0.4, 21, TEXT_COLOR)
    add_bullets(slide, ["No rewriting.", "No page reordering.", "No added claims."], 2.35, 20)

    image_path = make_image()
    try:
        slide = presentation.slides.add_slide(blank)
        add_title(slide, "Image Re-layout")
        add_text(slide, "Synthetic gradient image reused as a supplied asset.", 0.9, 1.55, 6.6, 0.75, 18, TEXT_COLOR)
        slide.shapes.add_picture(str(image_path), Inches(6.6), Inches(1.55), width=Inches(5.7), height=Inches(3.21))
    finally:
        image_path.unlink(missing_ok=True)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Table Fidelity")
    table = slide.shapes.add_table(3, 3, Inches(1.0), Inches(1.7), Inches(10.8), Inches(3.0)).table
    rows = [["Metric", "Q1", "Q2"], ["Visits", "120", "145"], ["Conversion", "3.2%", "3.8%"]]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.text_frame.paragraphs[0].font.size = Pt(18)
            cell.text_frame.paragraphs[0].font.bold = row_index == 0

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Chart Fidelity")
    add_text(slide, "Quarterly Throughput", 0.9, 1.45, 5.0, 0.35, 18, TEXT_COLOR)
    chart_data = CategoryChartData()
    chart_data.categories = ["North", "South", "West"]
    chart_data.add_series("Actual", (18, 24, 21))
    chart_data.add_series("Plan", (20, 20, 20))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.9), Inches(11.3), Inches(4.8), chart_data
    ).chart
    chart.has_legend = True
    chart.legend.include_in_layout = False

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Density Check")
    add_bullets(slide, ["One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten"], 1.45, 16)

    slide = presentation.slides.add_slide(blank)
    add_text(slide, "Pause", 0.9, 3.05, 4.0, 0.8, 38, TITLE_COLOR, True)
    add_text(slide, "A near-empty divider page is still preserved 1:1.", 0.95, 4.0, 8.5, 0.4, 18, MUTED_COLOR)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("src/test/resources/beautify-pptx/phase0-baseline.pptx"),
        help="fixture output path",
    )
    args = parser.parse_args()
    build_presentation(args.output)
    print(f"Wrote synthetic phase-0 fixture: {args.output}")


if __name__ == "__main__":
    main()
